"""
TOR AI — FastAPI Backend
เชื่อมต่อกับ OpenTyphoon API (https://opentyphoon.ai)
รัน: python backend.py
"""

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from pydantic import BaseModel
from typing import List
import os, json, io, re
from openai import OpenAI

# ── Config ──────────────────────────────────────────────────────
# API_KEY   = os.environ.get("TYPHOON_API_KEY", "YOUR_API_KEY_HERE")
API_KEY = os.environ.get("TYPHOON_API_KEY")
BASE_URL  = "https://api.opentyphoon.ai/v1"
MODEL     = "typhoon-v2.5-30b-a3b-instruct"
PORT      = 8000

# ── FastAPI ──────────────────────────────────────────────────────
app = FastAPI(title="TOR AI Backend", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],        # รับจาก file:// และทุก origin
    allow_methods=["*"],
    allow_headers=["*"],
)

client = OpenAI(api_key=API_KEY, base_url=BASE_URL)

# ── Models ───────────────────────────────────────────────────────
class GenerateRequest(BaseModel):
    prompt: str
    stream: bool = True
    temperature: float = 0.7
    max_tokens: int = 1500

class TorSection(BaseModel):
    num: int
    title: str
    content: str

class ExportRequest(BaseModel):
    title: str
    agency: str = "หน่วยงาน"
    budget: str = ""
    project_type: str = ""
    criteria: str = ""
    created_at: str = ""
    sections: List[TorSection]

# ── Routes ───────────────────────────────────────────────────────
@app.get("/")
def serve_frontend():
    """เสิร์ฟหน้า HTML หลัก"""
    html_path = os.path.join(os.path.dirname(__file__), "tor_generator.html")
    return FileResponse(html_path)

@app.post("/api/export/docx")
def export_docx(req: ExportRequest):
    """สร้างไฟล์ Word (.docx) จาก TOR ทั้ง 10 ข้อ"""
    try:
        from docx import Document as DocxDocument
        from docx.shared import Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise HTTPException(status_code=500, detail="กรุณาติดตั้ง python-docx ก่อน: pip install python-docx")

    try:
        doc = DocxDocument()

        # ── ตั้งค่าหน้ากระดาษ A4 มาตรฐานราชการไทย ──
        sec = doc.sections[0]
        sec.page_width    = Cm(21)
        sec.page_height   = Cm(29.7)
        sec.left_margin   = Cm(3)
        sec.right_margin  = Cm(2)
        sec.top_margin    = Cm(2.5)
        sec.bottom_margin = Cm(2.5)

        FONT = "TH Sarabun New"

        def set_font(run, size=16, bold=False):
            run.bold = bold
            run.font.size = Pt(size)
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.insert(0, rFonts)
            for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
                rFonts.set(qn(attr), FONT)

        def add_text(para, text, bold=False, size=16, align=None):
            run = para.add_run(text)
            set_font(run, size=size, bold=bold)
            if align:
                para.alignment = align
            return run

        # ── หัวเรื่อง ──
        p = doc.add_paragraph()
        add_text(p, "ร่างขอบเขตของงาน (Terms of Reference)", bold=True, size=20,
                 align=WD_ALIGN_PARAGRAPH.CENTER)
        p = doc.add_paragraph()
        add_text(p, req.title, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
        doc.add_paragraph()

        # ── ข้อมูลโครงการ ──
        meta = [
            ("อ้างอิง",        "หนังสือเวียน กค (กวจ) 0405.4/ว 159 ลงวันที่ 20 มีนาคม 2566"),
            ("หน่วยงาน",       req.agency),
            ("ประเภทงาน",      req.project_type),
            ("วงเงินงบประมาณ", f"{req.budget} บาท" if req.budget else "-"),
            ("หลักเกณฑ์",      req.criteria),
            ("วันที่สร้าง",     req.created_at),
        ]
        for label, value in meta:
            if value and value.strip():
                p = doc.add_paragraph()
                add_text(p, f"{label}: ", bold=True, size=16)
                add_text(p, value, bold=False, size=16)
        doc.add_paragraph()

        # ── 10 หัวข้อ TOR ──
        import re
        for s in req.sections:
            p = doc.add_paragraph()
            add_text(p, f"ข้อ {s.num}  {s.title}", bold=True, size=16)
            for line in s.content.strip().split("\n"):
                line = line.strip()
                if not line:
                    continue
                # ข้ามบรรทัดที่เป็นหัวข้อซ้ำ
                line_no_num = re.sub(r'^(ข้อ\s*)?\d+[\.\s]*', '', line).strip()
                title_words = set(re.sub(r'[/\-]', ' ', s.title).split())
                line_words  = set(re.sub(r'[/\-]', ' ', line_no_num).split())
                # ถ้าคำในบรรทัดทับซ้อนกับชื่อหัวข้อ >= 60% และสั้นกว่า 60 ตัว → ข้าม
                if title_words and line_words:
                    overlap = len(title_words & line_words) / len(line_words)
                    if overlap >= 0.6 and len(line_no_num) < 60:
                        continue
                # ลดช่องว่างซ้ำซ้อนในเนื้อหา
                line = re.sub(r'[ \t]+', ' ', line).strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Cm(1)
                p.paragraph_format.space_after = Pt(4)
                add_text(p, line, size=16)
            doc.add_paragraph()

        # ── หมายเหตุท้ายเอกสาร ──
        p = doc.add_paragraph()
        add_text(p, "หมายเหตุ: ", bold=True, size=14)
        add_text(p, "โครงร่างนี้เป็นแนวทางเบื้องต้น เจ้าหน้าที่ผู้รับผิดชอบต้องตรวจสอบและแก้ไขก่อนนำไปใช้งาน", size=14)

        # ── บันทึกและส่งกลับ ──
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        from urllib.parse import quote
        safe_title = req.title[:40].strip()
        encoded = quote(safe_title, safe='')
        filename_ascii = f"TOR_{encoded}.docx"
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=\"TOR.docx\"; filename*=UTF-8''{filename_ascii}"}
        )
    except Exception as e:
        print(f"[export error] {e}")
        raise HTTPException(status_code=500, detail=f"สร้างไฟล์ไม่สำเร็จ: {str(e)}")


@app.get("/api/health")
def health():
    """ตรวจสอบว่า backend + API key พร้อมใช้งาน"""
    if API_KEY == "YOUR_API_KEY_HERE":
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    return {"status": "ok", "model": MODEL, "provider": "OpenTyphoon"}


@app.post("/api/generate")
def generate(req: GenerateRequest):
    """รับ prompt แล้วส่งไปที่ OpenTyphoon API พร้อม streaming"""
    if API_KEY == "YOUR_API_KEY_HERE":
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")

    def stream_generator():
        try:
            stream = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {
                        "role": "system",
                        "content": (
                            "คุณคือผู้เชี่ยวชาญด้านการจัดซื้อจัดจ้างภาครัฐไทย "
                            "ที่มีความเชี่ยวชาญในการร่างขอบเขตของงาน (TOR) "
                            "ตามมาตรฐานหนังสือเวียน กค (กวจ) 0405.4/ว 159 "
                            "ของกรมบัญชีกลาง "
                            "ตอบเป็นภาษาราชการไทยที่ถูกต้องและเป็นทางการ "
                            "ห้ามระบุยี่ห้อหรือรุ่นของสินค้าโดยตรง ให้ใช้ Functional Specification แทน "
                            "เนื้อหาในแต่ละข้อต้องสอดคล้องและสัมพันธ์กัน "
                            "ตอบเฉพาะเนื้อหาที่ถามโดยตรง ไม่ต้องมีคำนำหรือคำลงท้ายที่ไม่จำเป็น "
                            "ห้ามใช้อักขระภาษาจีน ญี่ปุ่น เกาหลี หรืออักขระพิเศษจากภาษาอื่นโดยเด็ดขาด "
                            "ใช้ตัวเลขอารบิก (1, 2, 3, 4, 5) เท่านั้น ห้ามใช้เลขไทย (๑, ๒, ๓, ๔, ๕) ในทุกกรณี "
                            "เมื่อเขียนรายการข้อย่อย ให้เริ่มต้นที่ข้อ 1 เสมอ "
                            "ห้ามแปลเนื้อหาเป็นภาษาอังกฤษโดยเด็ดขาด ห้ามมี paragraph หรือบรรทัดที่เป็นภาษาอังกฤษล้วน "
                            "ตอบเป็นภาษาไทยเท่านั้น คำศัพท์เทคนิคอาจมีภาษาอังกฤษแทรกในวงเล็บได้เท่านั้น"
                        )
                    },
                    {"role": "user", "content": req.prompt}
                ],
                temperature=req.temperature,
                max_completion_tokens=req.max_tokens,
                top_p=0.6,
                frequency_penalty=0,
                stream=True,
            )
            token_count = 0
            for chunk in stream:
                delta = chunk.choices[0].delta
                if delta.content is not None and delta.content != "":
                    token_count += 1
                    yield json.dumps({"response": delta.content}) + "\n"
            print(f"[stream] ส่ง {token_count} tokens")
            yield json.dumps({"response": "", "done": True}) + "\n"  # sentinel
        except Exception as e:
            print(f"[stream error] {e}")
            yield json.dumps({"error": str(e)}) + "\n"

    return StreamingResponse(
        stream_generator(),
        media_type="application/x-ndjson",
        headers={"X-Accel-Buffering": "no", "Cache-Control": "no-cache"}
    )


@app.get("/api/models")
def list_models():
    """ดูว่า API key นี้ใช้ model อะไรได้บ้าง"""
    if API_KEY == "YOUR_API_KEY_HERE":
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    try:
        models = client.models.list()
        return {"models": [m.id for m in models.data]}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/test")
def test_api():
    """ทดสอบ OpenTyphoon API ตรงๆ แบบไม่ stream"""
    if API_KEY == "YOUR_API_KEY_HERE":
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    try:
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": "คุณคือผู้เชี่ยวชาญด้านการจัดซื้อจัดจ้างภาครัฐไทย"},
                {"role": "user", "content": "ตอบว่า 'ทดสอบสำเร็จ' เท่านั้น"}
            ],
            max_completion_tokens=20,
            temperature=0.6,
            top_p=0.6,
            stream=False,
        )
        content = resp.choices[0].message.content
        return {"status": "ok", "response": content, "model": MODEL}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/generate/simple")
def generate_simple(req: GenerateRequest):
    """สำหรับทดสอบ — ไม่ stream คืน text เต็มทันที"""
    if API_KEY == "YOUR_API_KEY_HERE":
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    try:
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": req.prompt}],
            temperature=req.temperature,
            max_tokens=req.max_tokens,
            stream=False,
        )
        return {"response": resp.choices[0].message.content}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def _extract_text_from_pdf(file_bytes: bytes) -> str:
    """อ่านข้อความจากไฟล์ PDF (bytes) ด้วย pypdf"""
    try:
        import pypdf
    except ImportError:
        raise HTTPException(status_code=500, detail="กรุณาติดตั้ง pypdf ก่อน: pip install pypdf")
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text


def _auto_censor_text(text: str) -> str:
    """เซนเซอร์ข้อมูลติดต่อ (เบอร์โทร/อีเมล/เว็บไซต์) อัตโนมัติก่อนส่งให้ AI"""
    if not text:
        return ""
    text = re.sub(r'\b\d{2,3}-\d{3}-\d{4}\b|\b\d{2,3}-\d{4}-\d{4}\b|\b\d{9,10}\b', "[PHONE_NUMBER_HIDDEN]", text)
    text = re.sub(r'[\w\.-]+@[\w\.-]+\.\w+', "[EMAIL_HIDDEN]", text)
    text = re.sub(r'https?://[^\s<>"]+|www\.[^\s<>"]+', "[WEBSITE_HIDDEN]", text)
    return text


@app.post("/api/analyze-spec")
async def analyze_spec(
    job_description: str = Form(...),
    files: List[UploadFile] = File(...),
):
    """วิเคราะห์เอกสารสเปค/ใบเสนอราคาจากหลายบริษัท แล้วสรุปเป็นร่างสเปคกลาง (ไม่ล็อกยี่ห้อ)"""
    if API_KEY == "YOUR_API_KEY_HERE" or not API_KEY:
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="กรุณาระบุลักษณะงานหรือวัตถุประสงค์ที่ต้องการ")
    if len(files) < 2:
        raise HTTPException(status_code=400, detail="กรุณาอัปโหลดเอกสารอย่างน้อย 2 บริษัท เพื่อให้ระบบหาจุดร่วมของสเปคกลางได้")

    try:
        all_companies_data_prompt = ""
        company_names = []
        for idx, f in enumerate(files):
            company_label = f"[COMPANY_{idx+1}]"
            company_names.append(f.filename or company_label)
            raw_bytes = await f.read()
            raw_text = _extract_text_from_pdf(raw_bytes)
            clean_text = _auto_censor_text(raw_text)
            all_companies_data_prompt += f"\n--- ข้อมูลสเปคของ {company_label} ---\n"
            all_companies_data_prompt += clean_text[:5000] + "\n"

        prompt = f"""คุณคือผู้เชี่ยวชาญด้านการตรวจรับและจัดทำคุณลักษณะเฉพาะ (TOR Specialist)
งานของคุณคือวิเคราะห์สเปคจากข้อเสนอที่ได้รับ ({len(files)} บริษัท) แล้วสรุปเป็น 'ร่างสเปคกลาง' ที่ถูกต้องตามหลักกฎหมายจัดซื้อจัดจ้าง คือ "ห้ามระบุชื่อยี่ห้อหรือรุ่นสินค้าเด็ดขาด" แต่ให้ใช้เกณฑ์ทางเทคนิคที่ทุกบริษัทสามารถหาของมาสู้กันได้

[ลักษณะงานที่ผู้ใช้ต้องการ]:
{job_description}

[ข้อมูลเอกสารสเปคของทุกบริษัท]:
{all_companies_data_prompt}

กรุณาตอบกลับเป็นภาษาไทย โดยใช้รูปแบบ Markdown ที่กระชับ เป็นข้อๆ และเข้าใจง่ายที่สุด ดังนี้:

1. ## 📊 ตารางสรุปเปรียบเทียบสเปค (ทำเป็นตารางสั้นๆ สรุปเฉพาะจุดสำคัญ)

2. ## 📋 ร่างสเปคกลาง (ข้อกำหนดขั้นต่ำที่โปร่งใสและแข่งขันได้จริง)
ห้ามระบุคำว่า Intel, AMD, NVIDIA, GeForce หรือชื่อยี่ห้อ/รุ่นใดๆ โดยเด็ดขาด ให้ใช้คำอธิบายเชิงคุณลักษณะแทน เช่น
- หน่วยประมวลผลกลาง: ระบุจำนวนคอร์/เธรด และความเร็วสัญญาณนาฬิกาขั้นต่ำ
- หน่วยประมวลผลกราฟิก: ระบุประเภท (รวม/แยก) และหน่วยความจำขั้นต่ำ
- รายการอื่นๆ ให้ระบุค่าขั้นต่ำที่อย่างน้อย 3 รายผ่านเกณฑ์

3. ## 💡 ความเห็นกรรมการ (สรุปสั้น 3 บรรทัด)
- สเปคกลางนี้พอมั้ยกับงาน
- จุดที่ควรระวังหรือควร upgrade เพิ่มเพื่อความคุ้มค่า (สรุปเป็นข้อสั้นๆ ห้ามยาว)
"""
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": "คุณคือผู้เชี่ยวชาญด้านการตรวจรับและจัดทำคุณลักษณะเฉพาะ (TOR Specialist) ของหน่วยงานราชการไทย ตอบเป็นภาษาไทยเท่านั้น"},
                {"role": "user", "content": prompt},
            ],
            temperature=0.6,
            max_tokens=3000,
            stream=False,
        )
        return {
            "result": resp.choices[0].message.content,
            "companies": company_names,
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"วิเคราะห์ไม่สำเร็จ: {str(e)}")


def _ocr_pdf_to_text(file_bytes: bytes, filename: str = "document.pdf") -> str:
    """ดึงข้อความจาก PDF ด้วย Typhoon OCR (ทีละหน้า) — fallback เป็น pypdf หากใช้ OCR ไม่ได้"""
    import tempfile

    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name

        try:
            from typhoon_ocr import ocr_document
            import pypdf as _pypdf

            reader = _pypdf.PdfReader(io.BytesIO(file_bytes))
            num_pages = len(reader.pages)
            pages_text = []
            for page_num in range(1, num_pages + 1):
                try:
                    md = ocr_document(pdf_or_image_path=tmp_path, page_num=page_num)
                    pages_text.append(md or "")
                except Exception as ocr_err:
                    pages_text.append(f"[OCR หน้า {page_num} ล้มเหลว: {ocr_err}]")
            return "\n\n".join(pages_text)
        except ImportError:
            # ไม่มี typhoon-ocr ติดตั้ง — fallback ไปใช้ pypdf อ่านข้อความตรง ๆ
            return _extract_text_from_pdf(file_bytes)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


@app.post("/api/audit-tor")
async def audit_tor(file: UploadFile = File(...)):
    """ตรวจสอบความครบถ้วนและความสอดคล้องของเอกสาร ToR (PDF) ตามมาตรฐาน ว.159"""
    if API_KEY == "YOUR_API_KEY_HERE" or not API_KEY:
        raise HTTPException(status_code=401, detail="ยังไม่ได้ตั้ง TYPHOON_API_KEY")
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="กรุณาอัปโหลดไฟล์ PDF เท่านั้น")

    try:
        raw_bytes = await file.read()
        if not raw_bytes:
            raise HTTPException(status_code=400, detail="ไฟล์ว่างเปล่า อ่านเนื้อหาไม่ได้")

        doc_text = _ocr_pdf_to_text(raw_bytes, file.filename)
        if not doc_text.strip():
            raise HTTPException(status_code=422, detail="ไม่สามารถดึงข้อความจากไฟล์ PDF ได้ — ไฟล์อาจเสียหายหรือเป็นภาพคุณภาพต่ำ")

        # จำกัดความยาวข้อความที่ส่งให้ AI เพื่อไม่ให้เกินโควตา
        doc_text_for_ai = doc_text[:18000]

        prompt = f"""คุณคือผู้ช่วย screen เบื้องต้น (decision support) สำหรับตรวจร่างขอบเขตของงาน (ToR) งานจัดซื้อจัดจ้างภาครัฐไทย
อ้างอิงมาตรฐานหนังสือเวียน กค (กวจ) 0405.4/ว 159 ลงวันที่ 20 มีนาคม 2566 ซึ่งกำหนดว่า ToR ต้องมีหัวข้อมาตรฐานครบ 10 ข้อ ได้แก่
1. ความเป็นมา/เหตุผลความจำเป็น
2. วัตถุประสงค์
3. คุณสมบัติผู้เสนอราคา/ผู้ยื่นข้อเสนอ
4. แบบรูปรายการ/คุณลักษณะเฉพาะของพัสดุที่จะจัดซื้อจัดจ้าง
5. ระยะเวลาดำเนินการ
6. ระยะเวลาส่งมอบ/สถานที่ส่งมอบ
7. วงเงินงบประมาณ/ราคากลาง
8. หลักเกณฑ์การพิจารณาคัดเลือกข้อเสนอ
9. การรับประกัน/บำรุงรักษา (ถ้ามี)
10. หน่วยงานที่รับผิดชอบ/ผู้ติดต่อ

**บทบาทสำคัญ**: คุณเป็นผู้ช่วย screen เบื้องต้นเท่านั้น ไม่ใช่ผู้ตัดสินอนุมัติแทนเจ้าหน้าที่ ทุกข้อที่แจ้งเตือนต้องมีเหตุผลกำกับ ห้ามฟันธงเด็ดขาด ใช้ถ้อยคำเชิงเสนอแนะ เช่น "ควรพิจารณา…", "เข้าข่ายอาจขัด…" และเรื่องที่มีผลทางกฎหมาย (เช่น ล็อคสเปค) ต้องให้เจ้าหน้าที่เป็นผู้ยืนยันเสมอ

กรุณาตรวจเอกสาร ToR ต่อไปนี้ใน 3 ด้าน:

**ด้าน 1 — ความครบถ้วนของหัวข้อ**: ตรวจว่ามีหัวข้อมาตรฐานครบทั้ง 10 ข้อหรือไม่ ข้อใดขาดหรือมีแต่ไม่สมบูรณ์ ให้ระบุสถานะ "มี / ไม่มี / มีแต่ไม่สมบูรณ์" พร้อมอ้างอิงข้อความในเอกสาร (ถ้ามี)

**ด้าน 2 — ความสอดคล้องภายในเอกสาร**: ตรวจความสอดคล้องของตัวเลขกับคำอ่าน (เช่น "4,000 บาท (ห้าพันบาท)" คือผิด เพราะ 4,000 ต้องอ่านว่า "สี่พันบาท"), ความสอดคล้องของวันที่/ระยะเวลา, งบประมาณรวมกับรายการย่อย, การอ้างอิงข้ามข้อ

**ด้าน 3 — ความถูกต้องตามระเบียบ/ความเป็นกลาง**: ตรวจการล็อคสเปค เช่น ระบุยี่ห้อ/รุ่น/Part No. เฉพาะเจาะจงโดยไม่มีคำว่า "หรือเทียบเท่า", กำหนดคุณลักษณะที่ชี้ไปยังผู้ขายรายเดียว (ในสเปคหนึ่งควรมีผู้ผลิต/ผู้ขายที่ผ่านเกณฑ์ได้ไม่น้อยกว่า 3 ราย), เงื่อนไขคุณสมบัติที่กีดกันการแข่งขันเกินจำเป็น — เมื่อพบให้เสนอแนวทางปรับเป็นกลาง (ระบุสมรรถนะเชิงหน้าที่ + เติม "หรือเทียบเท่า") แต่ต้องระบุชัดว่าให้เจ้าหน้าที่ยืนยัน

ตอบกลับเป็นภาษาไทย รูปแบบ Markdown ตามโครงสร้างนี้เท่านั้น:

# ผลประเมิน ToR

## สรุป
- ความครบถ้วน: [n/10] หัวข้อ
- จุดที่พบ: ผิดระเบียบ [n] · ควรแก้ไข [n] · ข้อเสนอแนะ [n]

## ตารางตรวจความครบถ้วน 10 หัวข้อ (ว.159)
| ข้อ | หัวข้อมาตรฐาน | สถานะ | หมายเหตุ |
|---|---|---|---|
(แสดงครบทั้ง 10 แถว)

## รายการที่ควรตรวจสอบเพิ่มเติม
สำหรับแต่ละจุดที่พบ ให้ระบุ:
- ระดับ: 🔴 ผิดระเบียบ / 🟡 ควรแก้ไข / 🔵 ข้อเสนอแนะ
- ตำแหน่ง: [ข้อ/หัวข้อในเอกสาร]
- ปัญหา: [อธิบายสั้นๆ]
- เหตุผล/อ้างอิง: [อ้างอิงระเบียบที่เกี่ยวข้อง — ถ้าไม่แน่ใจให้บอกว่าไม่แน่ใจ]
- ข้อเสนอแก้ไข: [แนวทาง]
- ความมั่นใจ: [สูง / ปานกลาง / ต่ำ]

## คำเตือน
ปิดท้ายด้วยข้อความว่า: ผลนี้เป็นการ screen เบื้องต้นโดย AI การตัดสินใจขั้นสุดท้ายเป็นของเจ้าหน้าที่ผู้รับผิดชอบ

[เนื้อหาเอกสาร ToR ที่ดึงได้]:
{doc_text_for_ai}
"""
        resp = client.chat.completions.create(
            model=MODEL,
            messages=[
                {"role": "system", "content": "คุณคือผู้ช่วย screen เบื้องต้นสำหรับตรวจร่างขอบเขตของงาน (ToR) ของหน่วยงานราชการไทย ตามมาตรฐาน ว.159 ตอบเป็นภาษาไทยเท่านั้น และห้ามฟันธงแทนเจ้าหน้าที่ในเรื่องที่มีผลทางกฎหมาย"},
                {"role": "user", "content": prompt},
            ],
            temperature=0.4,
            max_tokens=4000,
            stream=False,
        )
        return {
            "result": resp.choices[0].message.content,
            "filename": file.filename,
            "extracted_chars": len(doc_text),
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ตรวจสอบไม่สำเร็จ: {str(e)}")


# ── Run ──────────────────────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn, sys

    key = os.environ.get("TYPHOON_API_KEY", "")
    if not key or key == "YOUR_API_KEY_HERE":
        print("\n" + "="*55)
        print("  ⚠️  กรุณาตั้งค่า API Key ก่อนรัน backend")
        print("="*55)
        print("\n  วิธีที่ 1 — ตั้งค่าผ่าน environment variable:")
        print("    Windows CMD:")
        print("      set TYPHOON_API_KEY=sk-xxxxxxxxxxxx")
        print("      python backend.py")
        print("\n  วิธีที่ 2 — แก้ในไฟล์ backend.py บรรทัด:")
        print('      API_KEY = "sk-xxxxxxxxxxxx"')
        print("\n  สมัครรับ API Key ฟรีที่: https://opentyphoon.ai")
        print("="*55 + "\n")

        ans = input("มี API Key แล้ว? กรอกได้เลย (หรือ Enter เพื่อออก): ").strip()
        if ans:
            os.environ["TYPHOON_API_KEY"] = ans
            client = OpenAI(api_key=ans, base_url=BASE_URL)
            print(f"✅ ตั้งค่า key เรียบร้อย\n")
        else:
            sys.exit(0)

    print(f"\n🚀 TOR AI Backend รันที่ http://localhost:{PORT}")
    print(f"   Model  : {MODEL}")
    print(f"   Docs   : http://localhost:{PORT}/docs")
    print(f"   Health : http://localhost:{PORT}/api/health\n")
    uvicorn.run(app, host="0.0.0.0", port=PORT, log_level="info")
